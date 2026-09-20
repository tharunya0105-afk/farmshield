# -*- coding: utf-8 -*-
"""Geometry of old deck's overlay shapes (AnswerBox/FlowChip/Icon) on slides 2,3,4,6."""
from pptx import Presentation
from pptx.util import Emu

OLD = r"C:\Users\USER\Downloads\Divacoded_AAA09_Hackwell2_0_with_visuals (1).pptx"
prs = Presentation(OLD)
for idx in (1, 2, 3, 5):
    s = prs.slides[idx]
    print(f"\n=== OLD slide {idx+1} overlays ===")
    for sh in s.shapes:
        if sh.shape_id > 1000 or sh.name.startswith(("Flow", "Icon", "Answer")):
            t = ""
            if sh.has_text_frame:
                t = sh.text_frame.paragraphs[0].text[:40] if sh.text_frame.paragraphs else ""
            try:
                print(f"  {sh.name[:26]:<26} L={Emu(sh.left).inches:.2f} T={Emu(sh.top).inches:.2f} W={Emu(sh.width).inches:.2f} H={Emu(sh.height).inches:.2f}  {t}")
            except Exception:
                print(f"  {sh.name[:26]:<26} (no geom) {t}")
