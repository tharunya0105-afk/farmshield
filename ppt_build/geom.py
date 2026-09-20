# -*- coding: utf-8 -*-
"""Dump geometry of template slides; extract old deck's architecture + logo images."""
from pptx import Presentation
from pptx.util import Emu

TMPL = r"C:\Users\USER\Downloads\H20_PPT_Template (1).pptx"
OLD  = r"C:\Users\USER\Downloads\Divacoded_AAA09_Hackwell2_0_with_visuals (1).pptx"

prs = Presentation(TMPL)
for idx in (2, 3, 4, 5, 6):  # slides 3,4,5,6,7 (0-based index)
    s = prs.slides[idx]
    print(f"\n=== TEMPLATE slide {idx+1} ===")
    for sh in s.shapes:
        t = ""
        if sh.has_text_frame:
            t = sh.text_frame.paragraphs[0].text[:40] if sh.text_frame.paragraphs else ""
        print(f"  id={sh.shape_id:<4} {sh.name[:22]:<22} L={Emu(sh.left).inches:.2f} T={Emu(sh.top).inches:.2f} W={Emu(sh.width).inches:.2f} H={Emu(sh.height).inches:.2f}  {t}")

# extract old deck images: architecture (slide 5) + its icons
old = Presentation(OLD)
s5 = old.slides[4]
print("\n=== OLD slide 5 pictures ===")
for sh in s5.shapes:
    if sh.shape_type == 13:
        blob = sh.image.blob
        fn = f"old_{sh.shape_id}_{sh.name.replace(' ', '_')}.png"
        with open(fn, "wb") as f:
            f.write(blob)
        print(f"  saved {fn}  ({len(blob)//1024} KB)  pos L={Emu(sh.left).inches:.2f} T={Emu(sh.top).inches:.2f} W={Emu(sh.width).inches:.2f} H={Emu(sh.height).inches:.2f}")
