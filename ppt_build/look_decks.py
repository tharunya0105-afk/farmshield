# -*- coding: utf-8 -*-
"""Inspect the H2O template and the old Divacoded deck."""
from pptx import Presentation
from pptx.util import Emu

def describe(path, label):
    prs = Presentation(path)
    print(f"===== {label} =====")
    print("slide size:", prs.slide_width, "x", prs.slide_height,
          f"({Emu(prs.slide_width).inches:.2f}in x {Emu(prs.slide_height).inches:.2f}in)")
    print("slide count:", len(prs.slides))
    for i, s in enumerate(prs.slides):
        print(f"\n--- slide {i+1} (layout: {s.slide_layout.name}) ---")
        for sh in s.shapes:
            kind = sh.shape_type
            txt = ""
            if sh.has_text_frame:
                txt = " | ".join(p.text for p in sh.text_frame.paragraphs if p.text)[:90]
            print(f"  {sh.shape_id:>3} {str(kind):<28} name={sh.name[:24]:<24} {txt}")

describe(r"C:\Users\USER\Downloads\H20_PPT_Template (1).pptx", "TEMPLATE H2O")
print("\n\n")
describe(r"C:\Users\USER\Downloads\Divacoded_AAA09_Hackwell2_0_with_visuals (1).pptx", "OLD DIVACODED DECK")
