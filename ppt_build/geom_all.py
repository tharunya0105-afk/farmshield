# -*- coding: utf-8 -*-
"""Dump geometry of template slides 1-7."""
from pptx import Presentation
from pptx.util import Emu

TMPL = r"C:\Users\USER\Downloads\H20_PPT_Template (1).pptx"
prs = Presentation(TMPL)
for idx in range(min(7, len(prs.slides))):
    s = prs.slides[idx]
    print(f"\n=== TEMPLATE slide {idx+1} ===")
    for sh in s.shapes:
        t = ""
        if sh.has_text_frame and sh.text_frame.text.strip():
            t = sh.text_frame.text.strip().replace("\n", " | ")[:52]
        try:
            print(f"  L{Emu(sh.left).inches:6.2f} T{Emu(sh.top).inches:6.2f} "
                  f"W{Emu(sh.width).inches:6.2f} H{Emu(sh.height).inches:6.2f}  {sh.shape_type}  {t}")
        except Exception as e:
            print(f"  <no geom> {sh.shape_type}  {t}")
