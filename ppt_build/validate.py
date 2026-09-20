# -*- coding: utf-8 -*-
"""Validate the generated FarmShield deck."""
import sys
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
from pptx import Presentation

OUT = r"C:\Users\USER\Downloads\FarmShield_AAA09_Hackwell2_0.pptx"
prs = Presentation(OUT)
print("slides:", len(prs.slides))

all_text = []
for i, s in enumerate(prs.slides, 1):
    texts = []
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            texts.append(sh.text_frame.text.strip())
    joined = " || ".join(texts)
    all_text.append((i, joined))
    print(f"\n--- slide {i} ({len(texts)} text shapes) ---")
    print(joined[:400])

full = "\n".join(t for _, t in all_text)
print("\n=== CHECKS ===")
for bad in ["Divacoded", "divacoded", "Team Name ", "TEAM NAME ", "1. Team Lead", "PROBLEM STATEMENT NUMBER", "YOLOv8) + RL", "80% Less"]:
    print(f"stale '{bad}':", "FOUND!" if bad in full else "clean")
for good in ["FarmShield", "AAA-09", "Tharunya", "2.1", "14.2", "78.6%", "prototype", "Tamil", "Detect Early"]:
    print(f"has '{good}':", "yes" if good in full else "MISSING!")
