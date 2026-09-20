# -*- coding: utf-8 -*-
"""Dump geometry of template slides 1-2 (title + problem)."""
from pptx import Presentation
from pptx.util import Emu

TMPL = r"C:\Users\USER\Downloads\H20_PPT_Template (1).pptx"
prs = Presentation(TMPL)
for idx in (0, 1):
    s = prs.slides[idx]
    print(f"\n=== TEMPLATE slide {idx+1} ===")
    for sh in s.shapes:
        t = ""
        if sh.has_text_frame:
            t = sh.text_frame.paragraphs[0].text[:44] if sh.text_frame.paragraphs else ""
        print(f"  id={sh.shape_id:<4} {sh.name[:22]:<22} L={Emu(sh.left).inches:.2f} T={Emu(sh.top).inches:.2f} W={Emu(sh.width).inches:.2f} H={Emu(sh.height).inches:.2f}  {t}")
